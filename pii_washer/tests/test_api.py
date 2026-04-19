"""Integration tests for the PII Washer FastAPI wrapper (B2)."""

import io

import pytest
from fastapi.testclient import TestClient

from pii_washer.session_manager import SessionManager

# ---------------------------------------------------------------------------
# Mock Detection Engine (same pattern as test_session_manager.py)
# ---------------------------------------------------------------------------

class MockDetectionEngine:
    """Lightweight mock that returns predictable detections without spaCy."""

    def detect(self, text, confidence_threshold=0.3, language="en"):
        detections = []
        counter = 1
        test_patterns = [
            ("John Smith", "NAME"),
            ("Jane Doe", "NAME"),
            ("john@example.com", "EMAIL"),
            ("(555) 123-4567", "PHONE"),
            ("219-09-9999", "SSN"),
            ("Springfield", "ADDRESS"),
        ]
        for value, category in test_patterns:
            start = 0
            while True:
                pos = text.find(value, start)
                if pos == -1:
                    break
                detections.append({
                    "id": f"pii_{counter:03d}",
                    "category": category,
                    "original_value": value,
                    "positions": [{"start": pos, "end": pos + len(value)}],
                    "confidence": 0.85,
                })
                counter += 1
                start = pos + len(value)
        return detections

    def get_supported_categories(self):
        return ["NAME", "ADDRESS", "PHONE", "EMAIL", "SSN", "DOB", "CCN", "IP", "URL"]


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

SAMPLE_TEXT = "My name is John Smith and my email is john@example.com."
RESPONSE_TEXT = "Hello [Person_1], your email [Email_1] has been noted."


@pytest.fixture
def client():
    """TestClient backed by a fresh in-memory SessionManager with mock engine."""
    from pii_washer.api.main import create_app
    manager = SessionManager(detection_engine=MockDetectionEngine())
    app = create_app(session_manager=manager)
    with TestClient(app) as c:
        yield c


def _create_analyzed_session(client):
    """Helper: create and analyze a session, return (session_id, detection_id)."""
    r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
    session_id = r.json()["session_id"]
    r2 = client.post(f"/api/v1/sessions/{session_id}/analyze")
    detections = r2.json()["detections"]
    detection_id = detections[0]["id"] if detections else None
    return session_id, detection_id


# ---------------------------------------------------------------------------
# 1. Happy path — full workflow
# ---------------------------------------------------------------------------

class TestFullWorkflow:
    def test_end_to_end_workflow(self, client):
        # 1. Create session
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        assert r.status_code == 201
        session_id = r.json()["session_id"]

        # 2. Analyze
        r = client.post(f"/api/v1/sessions/{session_id}/analyze")
        assert r.status_code == 200
        detections = r.json()["detections"]
        assert len(detections) > 0

        # 3. Confirm one detection
        det_id = detections[0]["id"]
        r = client.patch(
            f"/api/v1/sessions/{session_id}/detections/{det_id}",
            json={"status": "confirmed"},
        )
        assert r.status_code == 200
        assert r.json()["status"] == "confirmed"

        # 4. Confirm all remaining
        r = client.post(f"/api/v1/sessions/{session_id}/detections/confirm-all")
        assert r.status_code == 200
        assert r.json()["confirmed_count"] >= 0

        # 5. Depersonalize
        r = client.post(f"/api/v1/sessions/{session_id}/depersonalize")
        assert r.status_code == 200
        body = r.json()
        assert "depersonalized_text" in body
        assert body["confirmed_count"] > 0

        # 6. Load response
        deperso_text = body["depersonalized_text"]
        r = client.post(
            f"/api/v1/sessions/{session_id}/response",
            json={"text": deperso_text},
        )
        assert r.status_code == 200
        assert r.json()["status"] == "awaiting_response"

        # 7. Repersonalize
        r = client.post(f"/api/v1/sessions/{session_id}/repersonalize")
        assert r.status_code == 200
        body = r.json()
        assert "repersonalized_text" in body
        assert "match_summary" in body

        # 8. Verify PII is restored
        repersonalized = body["repersonalized_text"]
        assert "John Smith" in repersonalized or "john@example.com" in repersonalized


# ---------------------------------------------------------------------------
# 2. Session CRUD
# ---------------------------------------------------------------------------

class TestSessionCRUD:
    def test_create_from_text_returns_201(self, client):
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        assert r.status_code == 201
        body = r.json()
        assert "session_id" in body
        assert body["status"] == "user_input"
        assert body["source_format"] == "paste"
        assert body["source_filename"] is None

    def test_create_from_file_upload_returns_201(self, client):
        content = b"Hello, my name is John Smith."
        r = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("test.txt", io.BytesIO(content), "text/plain")},
        )
        assert r.status_code == 201
        body = r.json()
        assert "session_id" in body
        assert body["source_format"] == ".txt"
        assert body["source_filename"] == "test.txt"

    def test_create_session_does_not_return_original_text(self, client):
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        assert r.status_code == 201
        assert "original_text" not in r.json()

    def test_get_session_by_id_returns_full_data(self, client):
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        session_id = r.json()["session_id"]
        r2 = client.get(f"/api/v1/sessions/{session_id}")
        assert r2.status_code == 200
        body = r2.json()
        assert body["session_id"] == session_id
        assert body["original_text"] == SAMPLE_TEXT
        assert "pii_detections" in body

    def test_get_session_status(self, client):
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        session_id = r.json()["session_id"]
        r2 = client.get(f"/api/v1/sessions/{session_id}/status")
        assert r2.status_code == 200
        body = r2.json()
        assert body["session_id"] == session_id
        assert "detection_count" in body
        assert "can_analyze" in body

    def test_get_session_returns_typed_response(self, client):
        """GET /sessions/{id} must return exactly the expected fields."""
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        session_id = r.json()["session_id"]
        r2 = client.get(f"/api/v1/sessions/{session_id}")
        assert r2.status_code == 200
        body = r2.json()
        expected_keys = {
            "session_id", "status", "created_at", "updated_at",
            "source_format", "source_filename", "original_text",
            "pii_detections", "depersonalized_text", "response_text",
            "repersonalized_text", "unmatched_placeholders",
        }
        assert set(body.keys()) == expected_keys


# ---------------------------------------------------------------------------
# 3. File upload validation
# ---------------------------------------------------------------------------

class TestFileUpload:
    def test_upload_txt_succeeds(self, client):
        content = b"Contact John Smith at john@example.com"
        r = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("doc.txt", io.BytesIO(content), "text/plain")},
        )
        assert r.status_code == 201

    def test_upload_md_succeeds(self, client):
        content = b"# Note\nCall (555) 123-4567 to reach John Smith."
        r = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("notes.md", io.BytesIO(content), "text/markdown")},
        )
        assert r.status_code == 201

    def test_upload_unsupported_extension_returns_422(self, client):
        content = b"print('hello')"
        r = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("script.py", io.BytesIO(content), "text/plain")},
        )
        assert r.status_code == 422
        assert r.json()["error"]["code"] == "UNSUPPORTED_FORMAT"

    def test_upload_metadata_persists_on_subsequent_reads(self, client):
        """Bug #1: source_format and source_filename must survive past the initial response."""
        content = b"Hello, my name is John Smith."
        r = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("report.txt", io.BytesIO(content), "text/plain")},
        )
        assert r.status_code == 201
        session_id = r.json()["session_id"]

        r2 = client.get(f"/api/v1/sessions/{session_id}")
        assert r2.status_code == 200
        body = r2.json()
        assert body["source_format"] == ".txt", f"Expected '.txt', got '{body['source_format']}'"
        assert body["source_filename"] == "report.txt", f"Expected 'report.txt', got '{body['source_filename']}'"

    def test_upload_exceeding_size_limit_returns_413(self, client):
        large_content = b"A" * (1_048_576 + 1)
        r = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("big.txt", io.BytesIO(large_content), "text/plain")},
        )
        assert r.status_code == 413
        assert r.json()["error"]["code"] == "FILE_TOO_LARGE"

    def test_upload_docx_file(self, client):
        from docx import Document
        doc = Document()
        doc.add_paragraph("John Smith lives in Springfield.")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        resp = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert resp.status_code == 201
        data = resp.json()
        assert data["source_format"] == ".docx"
        assert data["source_filename"] == "test.docx"

    def test_upload_pdf_file(self, client):
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        c.drawString(72, 750, "Jane Doe called (555) 123-4567 yesterday.")
        c.showPage()
        c.save()
        buf.seek(0)
        resp = client.post("/api/v1/sessions/upload", files={"file": ("test.pdf", buf, "application/pdf")})
        assert resp.status_code == 201
        assert resp.json()["source_format"] == ".pdf"

    def test_upload_csv_file(self, client):
        resp = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("test.csv", b"Name,Email\nJohn Smith,john@example.com\n", "text/csv")},
        )
        assert resp.status_code == 201
        assert resp.json()["source_format"] == ".csv"

    def test_upload_xlsx_file(self, client):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Email"])
        ws.append(["John Smith", "john@example.com"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = client.post(
            "/api/v1/sessions/upload",
            files={"file": ("test.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 201
        assert resp.json()["source_format"] == ".xlsx"

    def test_upload_html_file(self, client):
        content = b"<html><body><p>John Smith called (555) 123-4567.</p></body></html>"
        resp = client.post("/api/v1/sessions/upload",
            files={"file": ("test.html", content, "text/html")})
        assert resp.status_code == 201
        assert resp.json()["source_format"] == ".html"


# ---------------------------------------------------------------------------
# 4. Detection management
# ---------------------------------------------------------------------------

class TestDetectionManagement:
    def test_confirm_detection_changes_status(self, client):
        session_id, det_id = _create_analyzed_session(client)
        r = client.patch(
            f"/api/v1/sessions/{session_id}/detections/{det_id}",
            json={"status": "confirmed"},
        )
        assert r.status_code == 200
        assert r.json()["status"] == "confirmed"
        assert r.json()["id"] == det_id

    def test_reject_detection_changes_status(self, client):
        session_id, det_id = _create_analyzed_session(client)
        r = client.patch(
            f"/api/v1/sessions/{session_id}/detections/{det_id}",
            json={"status": "rejected"},
        )
        assert r.status_code == 200
        assert r.json()["status"] == "rejected"

    def test_edit_placeholder(self, client):
        session_id, det_id = _create_analyzed_session(client)
        r = client.patch(
            f"/api/v1/sessions/{session_id}/detections/{det_id}/placeholder",
            json={"placeholder": "[My_Client]"},
        )
        assert r.status_code == 200
        assert r.json()["placeholder"] == "[My_Client]"

    def test_confirm_all_returns_count(self, client):
        session_id, _ = _create_analyzed_session(client)
        r = client.post(f"/api/v1/sessions/{session_id}/detections/confirm-all")
        assert r.status_code == 200
        assert "confirmed_count" in r.json()
        assert r.json()["confirmed_count"] >= 0

    def test_manual_add_detection(self, client):
        text = "Contact Dr. Martinez about the referral."
        r = client.post("/api/v1/sessions", json={"text": text})
        session_id = r.json()["session_id"]
        client.post(f"/api/v1/sessions/{session_id}/analyze")

        r2 = client.post(
            f"/api/v1/sessions/{session_id}/detections",
            json={"text_value": "Dr. Martinez", "category": "NAME"},
        )
        assert r2.status_code == 201
        body = r2.json()
        assert body["original_value"] == "Dr. Martinez"
        assert body["category"] == "NAME"
        assert body["source"] == "manual"
        assert body["occurrences_found"] >= 1
        assert len(body["positions"]) >= 1


# ---------------------------------------------------------------------------
# 5. Error handling
# ---------------------------------------------------------------------------

class TestErrorHandling:
    def test_get_nonexistent_session_returns_404(self, client):
        r = client.get("/api/v1/sessions/nonexistent")
        assert r.status_code == 404
        assert r.json()["error"]["code"] == "NOT_FOUND"

    def test_analyze_wrong_state_returns_409(self, client):
        session_id, _ = _create_analyzed_session(client)
        # Already in 'analyzed' state — cannot analyze again
        r = client.post(f"/api/v1/sessions/{session_id}/analyze")
        assert r.status_code == 409
        assert r.json()["error"]["code"] == "INVALID_STATE"

    def test_depersonalize_no_confirmed_returns_422(self, client):
        session_id, _ = _create_analyzed_session(client)
        # No detections confirmed
        r = client.post(f"/api/v1/sessions/{session_id}/depersonalize")
        assert r.status_code == 422
        assert r.json()["error"]["code"] == "VALIDATION_ERROR"

    def test_update_detection_invalid_status_returns_422(self, client):
        session_id, det_id = _create_analyzed_session(client)
        r = client.patch(
            f"/api/v1/sessions/{session_id}/detections/{det_id}",
            json={"status": "not_a_real_status"},
        )
        assert r.status_code == 422
        assert r.json()["error"]["code"] == "VALIDATION_ERROR"

    def test_update_missing_detection_returns_404(self, client):
        session_id, _ = _create_analyzed_session(client)
        r = client.patch(
            f"/api/v1/sessions/{session_id}/detections/pii_999",
            json={"status": "confirmed"},
        )
        assert r.status_code == 404
        assert r.json()["error"]["code"] == "NOT_FOUND"

    def test_manual_add_invalid_category_returns_422(self, client):
        session_id, _ = _create_analyzed_session(client)
        r = client.post(
            f"/api/v1/sessions/{session_id}/detections",
            json={"text_value": "Something", "category": "INVALID_CATEGORY"},
        )
        assert r.status_code == 422
        assert r.json()["error"]["code"] == "VALIDATION_ERROR"

    def test_manual_add_duplicate_returns_409(self, client):
        session_id, _ = _create_analyzed_session(client)
        # "John Smith" is already detected as NAME by the mock engine
        r = client.post(
            f"/api/v1/sessions/{session_id}/detections",
            json={"text_value": "John Smith", "category": "NAME"},
        )
        assert r.status_code == 409
        assert r.json()["error"]["code"] == "DUPLICATE_DETECTION"

    def test_manual_add_text_not_found_returns_422(self, client):
        session_id, _ = _create_analyzed_session(client)
        r = client.post(
            f"/api/v1/sessions/{session_id}/detections",
            json={"text_value": "Completely Absent Text XYZ123", "category": "NAME"},
        )
        assert r.status_code == 422
        assert r.json()["error"]["code"] == "TEXT_NOT_FOUND"

    def test_server_error_does_not_leak_exception_details(self, client):
        """500 responses must not expose internal exception class or message."""
        r = client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        session_id = r.json()["session_id"]
        from unittest.mock import patch
        with patch.object(client.app.state.session_manager, "get_session", side_effect=TypeError("internal details")):
            r = client.get(f"/api/v1/sessions/{session_id}")
        assert r.status_code == 500
        msg = r.json()["error"]["message"]
        assert "TypeError" not in msg
        assert "internal details" not in msg
        assert msg == "An unexpected error occurred"


# ---------------------------------------------------------------------------
# 6. Health check
# ---------------------------------------------------------------------------

class TestHealthCheck:
    def test_health_returns_200(self, client):
        r = client.get("/api/v1/health")
        assert r.status_code == 200
        body = r.json()
        assert body["status"] == "ok"
        assert "engine_available" in body
        assert "version" in body

    def test_health_engine_available_with_mock(self, client):
        r = client.get("/api/v1/health")
        assert r.json()["engine_available"] is True

    def test_health_engine_unavailable_when_no_engine(self):
        from pii_washer.api.main import create_app
        manager = SessionManager(detection_engine=None)
        app = create_app(session_manager=manager)
        with TestClient(app) as c:
            r = c.get("/api/v1/health")
            assert r.status_code == 200
            assert r.json()["engine_available"] is False

    def test_version_matches_package_metadata(self, client):
        """Version reported by the API must match the installed package metadata.

        This guards the single-source-of-truth contract: pyproject.toml is the
        authority, and the runtime reads it via importlib.metadata.
        """
        from importlib.metadata import version as pkg_version

        expected = pkg_version("pii-washer")

        health = client.get("/api/v1/health").json()
        assert health["version"] == expected

        # FastAPI exposes the app-level version via openapi.json
        openapi = client.get("/openapi.json").json()
        assert openapi["info"]["version"] == expected


# ---------------------------------------------------------------------------
# 7. Session lifecycle cleanup
# ---------------------------------------------------------------------------

class TestReset:
    def test_reset_clears_session(self, client):
        client.post("/api/v1/sessions", json={"text": SAMPLE_TEXT})
        r = client.post("/api/v1/sessions/reset")
        assert r.status_code == 200
        assert r.json()["deleted_count"] >= 1

    def test_reset_when_empty(self, client):
        r = client.post("/api/v1/sessions/reset")
        assert r.status_code == 200
        assert r.json()["deleted_count"] == 0


